"""mystery.apps.reports.excel_report — AnalysisResult 列表 → Excel。

函数签名（002.md W1-A）::

    def write_excel(results: list[dict], path: str) -> str

results = ``[AnalysisResult.to_dict(), ...]``。生成器内禁止取数/调 analyze。

W14：取消个股详情 sheet，全部个股详情以列补充进「汇总报告」单页；
新增年线过滤系统各条件达成情况列（来自 ``mystery.signal['年线条件']``，
analyze 阶段由 basic_filter 逐项判定写入，报表只读不改判）。
"""
from __future__ import annotations

import io
import logging
from typing import Any, Dict, List, Optional

import pandas as pd

logger = logging.getLogger(__name__)

SUMMARY_COLS = [
    ("symbol", "代码"),
    ("name", "名称"),
    ("score", "综合评分"),
    ("advice", "操作建议"),
    ("true_resonance", "真三振"),
    ("sector_name", "行业"),
    ("price", "最新价"),
    ("main_wave", "主升浪状态"),
    ("checklist_n", "主升浪满足"),
    ("platform", "平台状态"),
    ("vap_upper", "VAP-ATR上轨"),
    ("vap_break", "VAP-ATR突破"),
    ("pe", "PE"),
    ("pb", "PB"),
    ("trade_date", "分析日期"),
]

# 年线过滤系统各条件（与 basic_filter._basic_filter_checks 键一一对应，
# 顺序即展示顺序；值为 ✅/❌/-）。
YEARLINE_COND_KEYS = [
    "收盘价>MA250", "收盘价>MA60", "均线多头排列",
    "MA5>MA250", "MA10>MA250", "MA20>MA250", "MA60>MA250",
]

# 个股详情补充列：与 SUMMARY_COLS 完全同义的项不重复展示（值不同义的保留）。
_DETAIL_SKIP = {
    "股票代码", "股票名称", "分析日期", "最新价", "综合评分", "操作建议",
    "真三振", "行业", "主升浪状态", "主升浪满足", "平台状态", "PE", "PB",
}
# section 前缀（避免「突破信号」等跨段同名冲突）
_SECTION_PREFIX = {
    "基础信息": "",
    "三振共振": "三振.",
    "主升浪": "主升浪.",
    "平台突破": "平台.",
    "自适应 VAP-ATR": "VAP.",
    "财务": "财务.",
    "缠论结构": "缠论.",
}


def _flag(v) -> str:
    """布尔 → ✅/❌（None/缺失 → '-'）。"""
    if v is None:
        return "-"
    return "✅" if v else "❌"


def _flat(d: Dict[str, Any]) -> Dict[str, Any]:
    """把嵌套 to_dict() 拍平成汇总行（只读已有字段，不重算）。"""
    m = d.get("mystery", {}) or {}
    vap = m.get("vap_atr", {}) or {}
    plat = m.get("platform", {}) or {}
    main_wave = m.get("main_wave", {}) or {}
    cl = m.get("checklist8", {}) or {}
    fin = d.get("financial", {}) or {}
    sec = d.get("sector", {}) or {}
    pr = vap.get("平台范围") or {}
    row: Dict[str, Any] = {
        "symbol": d.get("symbol", ""),
        "name": d.get("name") or "未知",
        "score": d.get("score"),
        "advice": d.get("advice", ""),
        "true_resonance": _flag(d.get("true_resonance")),
        "sector_name": sec.get("行业名称", "未知"),
        "price": d.get("price"),
        "main_wave": main_wave.get("主升浪状态", "-"),
        "checklist_n": f"{cl.get('满足数量', 0)}/8" if cl else "-",
        "platform": plat.get("平台状态", "-"),
        "vap_upper": vap.get("自适应上轨") or pr.get("上沿"),
        "vap_break": "✅" if vap.get("突破信号") or plat.get("突破信号") else "❌",
        "pe": fin.get("PE"),
        "pb": fin.get("PB"),
        "trade_date": d.get("trade_date", ""),
    }
    # 年线过滤系统各条件达成情况（W14）
    sig = m.get("signal", {}) or {}
    yl = sig.get("年线条件") or {}
    for key in YEARLINE_COND_KEYS:
        row[f"yl:{key}"] = _flag(yl.get(key))
    # 个股详情补充列（W14：原个股详情 sheet 内容并入汇总页）
    row.update(_detail_flat(d))
    return row


def _detail_flat(d: Dict[str, Any]) -> Dict[str, Any]:
    """个股详情键值对 → 扁平列 dict（只读已有字段，不重算）。

    跳过 section 标题/空行/与 SUMMARY_COLS 同义重复项；列名带 section
    前缀防跨段同名冲突。
    """
    m = d.get("mystery", {}) or {}
    vap = m.get("vap_atr", {}) or {}
    plat = m.get("platform", {}) or {}
    main_wave = m.get("main_wave", {}) or {}
    res = m.get("resonance", {}) or {}
    sig = m.get("signal", {}) or {}
    fin = d.get("financial", {}) or {}
    sec = d.get("sector", {}) or {}
    rows: List[tuple] = []
    section = ""

    def add(title: str, val: Any):
        rows.append((section, title, val))

    add("股票代码", d.get("symbol", ""))
    add("股票名称", d.get("name") or "未知")
    add("分析日期", d.get("trade_date", ""))
    add("最新价", d.get("price"))
    add("综合评分", d.get("score"))
    add("操作建议", d.get("advice", ""))
    add("真三振", _flag(d.get("true_resonance")))
    add("行业", sec.get("行业名称", "未知"))
    add("行业趋势分", sec.get("行业趋势分"))

    def sec_rows(title: str, items: List[tuple]):
        nonlocal section
        section = title
        for k, v in items:
            rows.append((section, k, v))

    sec_rows("三振共振", [
        ("个股趋势", _flag(res.get("个股趋势"))),
        ("行业趋势", _flag(res.get("行业趋势"))),
        ("大盘趋势", _flag(res.get("大盘趋势"))),
        ("共振评分", res.get("共振评分")),
        ("共振级别", res.get("共振级别", "-")),
        ("年线滤网", _flag(sig.get("年线滤网"))),
    ])
    for i, line in enumerate((res.get("详情") or [])[:6], 1):
        rows.append((section, f"共振详情{i}", line))

    sec_rows("主升浪", [
        ("主升浪状态", main_wave.get("主升浪状态", "-")),
        ("持股状态", _flag(main_wave.get("持股状态"))),
        ("空中加油", _flag(main_wave.get("空中加油"))),
        ("MA5斜率", main_wave.get("MA5斜率")),
    ])
    cl = m.get("checklist8", {}) or {}
    rows.append((section, "主升浪满足", f"{cl.get('满足数量', 0)}/8" if cl else "-"))
    rows.append((section, "主升浪综合判断", cl.get("综合判断", "-")))
    for i, line in enumerate((main_wave.get("判定依据") or [])[:5], 1):
        rows.append((section, f"判定依据{i}", line))

    sec_rows("平台突破", [
        ("平台状态", plat.get("平台状态", "-")),
        ("突破信号", _flag(plat.get("突破信号"))),
        ("买横信号", _flag(plat.get("买横信号"))),
    ])
    pr = plat.get("平台范围") or {}
    if pr:
        rows.append((section, "平台箱体", f"{pr.get('下沿')} ~ {pr.get('上沿')}"))
    fixed = plat.get("固定箱体") or {}
    if fixed:
        rows.append((section, "固定箱体(近20日)", f"{fixed.get('下沿')} ~ {fixed.get('上沿')}"))
    rows.append((section, "多周期箱体状态", plat.get("多周期箱体状态", "-")))

    sec_rows("自适应 VAP-ATR", [
        ("POC(筹码控制点)", vap.get("POC")),
        ("自适应上轨", vap.get("自适应上轨")),
        ("自适应下轨", vap.get("自适应下轨")),
        ("ATR", vap.get("ATR")),
        ("突破信号", _flag(vap.get("突破信号"))),
    ])
    cyc = vap.get("自适应周期") or {}
    if cyc:
        rows.append((section, "自适应周期",
                     f"N={cyc.get('adaptive_n')} 快ATR={cyc.get('atr_m')} k={cyc.get('k')}"))
        if cyc.get("avg_turnover") is not None:
            rows.append((section, "近20日均换手", f"{cyc.get('avg_turnover')}%"))
    for i, line in enumerate((vap.get("详情") or [])[:4], 1):
        rows.append((section, f"VAP详情{i}", line))

    sec_rows("财务", [
        ("PE", fin.get("PE")),
        ("PB", fin.get("PB")),
        ("ROE", fin.get("roe")),
        ("股息", fin.get("divid_cash")),
        ("报告期", fin.get("report_date", "")),
    ])

    # 缠论摘要列（W3-A：chan 关时为空，不改变 score）
    chan = d.get("chan", {}) or {}
    if chan:
        c1 = chan.get("1d") or {}
        cw = chan.get("1w") or {}
        sec_rows("缠论结构", [
            ("日线末笔方向", "向上" if c1.get("last_bi_dir") == "up" else (
                "向下" if c1.get("last_bi_dir") == "down" else "-")),
            ("日线末笔确认", _flag(c1.get("last_bi_confirmed"))),
            ("日线在中枢内", _flag(c1.get("in_zs"))),
            ("周线末笔方向", "向上" if cw.get("last_bi_dir") == "up" else (
                "向下" if cw.get("last_bi_dir") == "down" else "-")),
            ("czsc 版本", d.get("czsc_ver", "")),
        ])

    out: Dict[str, Any] = {}
    for sec_name, item, val in rows:
        if item in _DETAIL_SKIP:
            continue
        col = f"{_SECTION_PREFIX.get(sec_name, '')}{item}"
        out[col] = val
    return out


def _build_buffer(results: List[Dict[str, Any]]) -> "io.BytesIO":
    """生成 Excel 到内存 buffer（单「汇总报告」页，含全部个股详情列）。

    W14：取消个股详情 sheet 与代码列超链接/导航（无目标 sheet）；
    汇总报告 = 核心列 + 年线条件列 + 个股详情补充列。
    write_excel 与 excel_bytes 共用，保证 web 下载与 daily 落盘格式一致。
    """
    rows = [_flat(d) for d in results]
    if rows:
        rows.sort(key=lambda r: (r["score"] is not None, float(r["score"] or -1)),
                  reverse=True)
    df = pd.DataFrame(rows)
    # 固定列顺序：核心列 + 年线条件列 + 详情列（按出现顺序）
    col_order = [k for k, _ in SUMMARY_COLS]
    col_order += [f"yl:{k}" for k in YEARLINE_COND_KEYS]
    for k in df.columns:
        if k not in col_order:
            col_order.append(k)
    df = df[col_order]
    rename = dict(SUMMARY_COLS)
    rename.update({f"yl:{k}": f"年线:{k}" for k in YEARLINE_COND_KEYS})
    df = df.rename(columns=rename)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="汇总报告", index=False)
        ws = writer.sheets["汇总报告"]
        # 表头加粗 + 冻结首行，列宽取 表头/值 长度（上限 26，避免明细列无限撑宽）
        from openpyxl.styles import Font
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col_idx, col in enumerate(df.columns, 1):
            letter = ws.cell(row=1, column=col_idx).column_letter
            max_len = max([len(str(col))] +
                          [len(str(v)) for v in df[col].tolist() if v is not None])
            ws.column_dimensions[letter].width = min(max(max_len, 8), 26)
    buf.seek(0)
    return buf


def excel_bytes(results: List[Dict[str, Any]]) -> bytes:
    """Web 下载用：AnalysisResult 列表 → xlsx bytes（格式同 write_excel）。"""
    return _build_buffer(results).getvalue()


def write_excel(results: List[Dict[str, Any]], path: str,
                qa_line: Optional[str] = None) -> str:
    """单「汇总报告」页（含全部个股详情列 + 年线条件列）。返回写入路径。

    qa_line：W26b 换手覆盖率 QA，写入表格下方一行（只读观测，不加列、不改排序）。
    """
    import os
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    buf = _build_buffer(results)
    with open(path, "wb") as f:
        f.write(buf.getvalue())
    if qa_line:
        _append_qa_row(path, qa_line)
    logger.info("✅ Excel 报告生成完成: %s（%d 只，单汇总页）", path, len(results))
    return path


def _append_qa_row(path: str, qa_line: str) -> None:
    """在「汇总报告」sheet 表格末行下方空一行写 QA（openpyxl 追加）。"""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb["汇总报告"]
        row = ws.max_row + 2
        ws.cell(row=row, column=1, value=qa_line)
        wb.save(path)
    except Exception as e:  # noqa: BLE001 QA 观测失败不阻塞报告
        logger.warning("Excel QA 行写入失败（不影响报告主体）: %s", e)
