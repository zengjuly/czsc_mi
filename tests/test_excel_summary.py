"""excel_report 单汇总页回归（W14）。

验证：
- 输出只有「汇总报告」一个 sheet（取消个股详情 sheet）
- 汇总页含个股详情补充列（三振/主升浪/平台/VAP/财务/缠论）
- 汇总页含年线过滤系统各条件达成情况列（来自 signal['年线条件']，值 ✅/❌）
- 代码列无超链接（无目标 sheet 可跳）
"""
from __future__ import annotations

import io

import openpyxl
import pytest

from mystery.apps.reports.excel_report import excel_bytes

_YEARLINE = {
    "收盘价>MA250": True, "收盘价>MA60": True, "均线多头排列": True,
    "MA5>MA250": True, "MA10>MA250": True, "MA20>MA250": True,
    "MA60>MA250": False,
}


def _mk(symbol: str, name: str, score: float) -> dict:
    return {
        "symbol": symbol, "name": name, "trade_date": "2026-08-30",
        "score": score, "advice": "持有", "true_resonance": True, "price": 10.0,
        "mystery": {"vap_atr": {"POC": 9.0, "突破信号": True, "自适应周期": {}},
                    "platform": {"平台状态": "平台", "突破信号": False,
                                 "买横信号": True, "平台范围": {"下沿": 8.0, "上沿": 12.0}},
                    "main_wave": {"主升浪状态": "主升浪", "持股状态": True,
                                  "空中加油": False, "MA5斜率": 0.3},
                    "resonance": {"个股趋势": True, "行业趋势": True,
                                  "大盘趋势": False, "共振评分": 60,
                                  "共振级别": "一级共振", "详情": ["a", "b"]},
                    "checklist8": {"满足数量": 6, "综合判断": "偏强"},
                    "signal": {"年线滤网": False, "年线条件": _YEARLINE},
                    "technical": {}},
        "financial": {"PE": 20.0, "PB": 2.5, "roe": 10.2, "divid_cash": 1.0,
                      "report_date": "2026-06"},
        "sector": {"行业名称": "测试", "行业趋势分": 15.0}, "chan": {},
        "czsc_ver": "",
    }


@pytest.fixture()
def wb():
    rs = [_mk("sh600519", "贵州茅台", 49.0),
          _mk("sz000001", "平安银行", 70.0),
          _mk("sh600150", "中国船舶", 20.0)]
    data = excel_bytes(rs)
    return openpyxl.load_workbook(io.BytesIO(data))


def test_single_summary_sheet(wb):
    """只有「汇总报告」一个 sheet，无个股详情 sheet。"""
    assert wb.sheetnames == ["汇总报告"]


def test_no_code_hyperlink(wb):
    """代码列无超链接（取消个股详情 sheet 后无跳转目标）。"""
    ws = wb["汇总报告"]
    for row in ws.iter_rows(min_row=2, max_col=1):
        for c in row:
            assert c.hyperlink is None


def test_yearline_condition_columns(wb):
    """年线过滤系统各条件达成情况列存在且值为 ✅/❌。"""
    ws = wb["汇总报告"]
    headers = [c.value for c in ws[1]]
    expect = [f"年线:{k}" for k in _YEARLINE]
    for col in expect:
        assert col in headers, f"缺年线条件列 {col}"
    # MA60>MA250 未达成 → 该列所有行应为 ❌
    idx = headers.index("年线:MA60>MA250") + 1
    for row in ws.iter_rows(min_row=2, max_col=idx):
        c = row[idx - 1]
        assert c.value == "❌", f"{c.coordinate} 应为 ❌ 实际 {c.value}"
    # 达成项 → ✅
    idx_ok = headers.index("年线:收盘价>MA250") + 1
    for row in ws.iter_rows(min_row=2, max_col=idx_ok):
        c = row[idx_ok - 1]
        assert c.value == "✅"


def test_detail_columns_present(wb):
    """个股详情补充列并入汇总页。"""
    ws = wb["汇总报告"]
    headers = [c.value for c in ws[1]]
    expect = ["三振.个股趋势", "三振.共振评分", "主升浪.持股状态",
              "主升浪.空中加油", "平台.买横信号", "平台.平台箱体",
              "VAP.POC(筹码控制点)", "财务.ROE", "财务.股息",
              "行业趋势分", "主升浪.主升浪综合判断"]
    for col in expect:
        assert col in headers, f"缺详情列 {col}"
    # 与 SUMMARY_COLS 同义项在详情补充列中不重复（裸列名只出现一次）
    assert headers.count("综合评分") == 1
    assert headers.count("操作建议") == 1


def test_detail_values_filled(wb):
    """详情列值真实落到行（取平安银行行验证）。"""
    ws = wb["汇总报告"]
    headers = [c.value for c in ws[1]]
    row_by_sym = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_by_sym[row[0]] = row
    r = row_by_sym["sz000001"]
    vals = dict(zip(headers, r))
    assert vals["三振.个股趋势"] == "✅"
    assert vals["主升浪.持股状态"] == "✅"
    assert vals["平台.买横信号"] == "✅"
    assert vals["平台.平台箱体"] == "8.0 ~ 12.0"
    assert vals["财务.ROE"] == 10.2
    assert vals["行业趋势分"] == 15.0


def test_sort_by_score_desc(wb):
    """按综合评分降序。"""
    ws = wb["汇总报告"]
    scores = [row[2] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert scores == sorted(scores, reverse=True)
